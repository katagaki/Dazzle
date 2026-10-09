"""Writes the sample presentation the App Store screenshots open.

Usage: python3 samples.py <en|ja>

Prints the directory it wrote the presentation to.

One deck serves both the iPhone and the iPad, since a slide scales to fit
either. Everything is set in Kivotos. Needs python-pptx; capture.sh runs this
through `uv run --with python-pptx`.
"""
import os, re, sys, tempfile, zipfile
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# The language is picked from a fixed list rather than taken as given, and the
# deck goes into a fresh private directory of the script's own making, whose
# path it prints for capture.sh to copy from and remove.
def choose(value, options):
    for option in options:
        if option == value:
            return option
    sys.exit(__doc__)


if len(sys.argv) != 2:
    sys.exit(__doc__)
LANG = choose(sys.argv[1], ("en", "ja"))
OUT = tempfile.mkdtemp(prefix=f"dazzle-samples-{LANG}-")
JA = LANG == "ja"


def t(en, ja):
    return ja if JA else en


# MARK: - Styles

# Both are on every iPhone and iPad, so the slides look as designed.
FONT = "Hiragino Sans" if JA else "Avenir Next"

# Schale's blues, from deep navy to ice.
INK = RGBColor(0x13, 0x25, 0x4A)
MUTED = RGBColor(0x5A, 0x6B, 0x8C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAPER = RGBColor(0xF3, 0xF8, 0xFF)
NAVY = RGBColor(0x0B, 0x3D, 0x91)
DEEP = RGBColor(0x0B, 0x5E, 0xD7)
BLUE = RGBColor(0x12, 0x8C, 0xFF)
CYAN = RGBColor(0x2F, 0xC6, 0xF6)
AZURE = RGBColor(0x6C, 0xC4, 0xFF)
ICE = RGBColor(0xBF, 0xE6, 0xFF)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
W, H = prs.slide_width, prs.slide_height


def solid(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def gradient(shape, top, bottom, angle=90):
    shape.fill.gradient()
    shape.fill.gradient_angle = angle
    stops = shape.fill.gradient_stops
    stops[0].color.rgb = top
    stops[0].position = 0
    stops[1].color.rgb = bottom
    stops[1].position = 1


def no_line(shape):
    shape.line.fill.background()


def shadow(shape, blur=Pt(18), distance=Pt(4), alpha=22):
    """A soft drop shadow, as an a:outerShdw on the shape."""
    spPr = shape._element.spPr
    effects = etree.SubElement(spPr, qn("a:effectLst"))
    outer = etree.SubElement(effects, qn("a:outerShdw"), blurRad=str(blur), dist=str(distance),
                             dir="5400000", algn="ctr", rotWithShape="0")
    color = etree.SubElement(outer, qn("a:srgbClr"), val="000000")
    etree.SubElement(color, qn("a:alpha"), val=str(alpha * 1000))


def background(slide, top, bottom=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    if bottom is None:
        solid(shape, top)
    else:
        gradient(shape, top, bottom, angle=135)
    no_line(shape)
    return shape


def text(slide, x, y, w, h, lines, size, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=None):
    """A text box with one paragraph per line; a line may be (text, size, color, bold)."""
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    for i, line in enumerate(lines if isinstance(lines, list) else [lines]):
        content, line_size, line_color, line_bold = (line + (None,) * 3)[:4] if isinstance(line, tuple) else (line, None, None, None)
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.alignment = align
        if spacing:
            paragraph.space_after = spacing
        run = paragraph.add_run()
        run.text = content
        run.font.name = FONT
        run.font.size = Pt(line_size or size)
        run.font.bold = bold if line_bold is None else line_bold
        run.font.color.rgb = line_color or color
    return box


def fill_text(shape, content, size, color=WHITE, bold=True, align=PP_ALIGN.CENTER):
    frame = shape.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = content
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def heading(slide, title, kicker, color=BLUE):
    bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.62), Inches(0.5), Inches(0.12))
    solid(bar, color)
    no_line(bar)
    text(slide, Inches(0.8), Inches(0.85), Inches(11.5), Inches(0.4), kicker, 16, color, bold=True)
    text(slide, Inches(0.8), Inches(1.25), Inches(11.5), Inches(0.9), title, 36, INK, bold=True)


def notes(slide, content):
    slide.notes_slide.notes_text_frame.text = content


# MARK: - Animations

def animate(slide, steps):
    """Gives the slide a main sequence that plays `steps` one click at a time.

    Each step is a list of (shape, effect) that play together, where effect is
    "fade" or "fly". Written the way PowerPoint writes these, so Dazzle reads
    them as the Fade and Fly In it offers.
    """
    ids = iter(range(3, 1000))
    groups = []
    for step in steps:
        effects = []
        for index, (shape, effect) in enumerate(step):
            spid = shape.shape_id
            node_type = "clickEffect" if index == 0 else "withEffect"
            target = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
            show = (f'<p:set><p:cBhvr><p:cTn id="{next(ids)}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/>'
                    f'</p:stCondLst></p:cTn>{target}<p:attrNameLst><p:attrName>style.visibility</p:attrName>'
                    f'</p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>')
            if effect == "fly":
                preset, subtype = 2, 4
                behaviours = show + "".join(
                    f'<p:anim calcmode="lin" valueType="num"><p:cBhvr additive="base"><p:cTn id="{next(ids)}" dur="500" '
                    f'fill="hold"/>{target}<p:attrNameLst><p:attrName>{name}</p:attrName></p:attrNameLst></p:cBhvr>'
                    f'<p:tavLst><p:tav tm="0"><p:val><p:strVal val="{start}"/></p:val></p:tav><p:tav tm="100000">'
                    f'<p:val><p:strVal val="#{name}"/></p:val></p:tav></p:tavLst></p:anim>'
                    for name, start in (("ppt_x", "#ppt_x"), ("ppt_y", "1+#ppt_h/2")))
            else:
                preset, subtype = 10, 0
                behaviours = show + (f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{next(ids)}" '
                                     f'dur="500"/>{target}</p:cBhvr></p:animEffect>')
            effects.append(f'<p:par><p:cTn id="{next(ids)}" presetID="{preset}" presetClass="entr" '
                           f'presetSubtype="{subtype}" fill="hold" grpId="0" nodeType="{node_type}"><p:stCondLst>'
                           f'<p:cond delay="0"/></p:stCondLst><p:childTnLst>{behaviours}</p:childTnLst></p:cTn></p:par>')
        groups.append(f'<p:par><p:cTn id="{next(ids)}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/>'
                      f'</p:stCondLst><p:childTnLst><p:par><p:cTn id="{next(ids)}" fill="hold"><p:stCondLst>'
                      f'<p:cond delay="0"/></p:stCondLst><p:childTnLst>{"".join(effects)}</p:childTnLst></p:cTn>'
                      f'</p:par></p:childTnLst></p:cTn></p:par>')
    builds = "".join(f'<p:bldP spid="{shape.shape_id}" grpId="0" animBg="1"/>'
                     for step in steps for shape, _ in step if shape.has_text_frame)
    xml = (f'<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
           f'<p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
           f'{"".join(groups)}</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl>'
           f'<p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0">'
           f'<p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par>'
           f'</p:tnLst>{"<p:bldLst>" + builds + "</p:bldLst>" if builds else ""}</p:timing>')
    slide._element.append(etree.fromstring(xml))


# MARK: - Slides

def title_slide():
    slide = prs.slides.add_slide(BLANK)
    background(slide, RGBColor(0x3A, 0xB4, 0xFF), DEEP)
    # A halo, as every student in Kivotos wears one.
    halo = slide.shapes.add_shape(MSO_SHAPE.DONUT, Inches(8.4), Inches(-1.6), Inches(6.4), Inches(6.4))
    halo.adjustments[0] = 0.07
    solid(halo, WHITE)
    no_line(halo)
    ring = slide.shapes.add_shape(MSO_SHAPE.DONUT, Inches(9.6), Inches(-0.4), Inches(4.0), Inches(4.0))
    ring.adjustments[0] = 0.12
    solid(ring, ICE)
    no_line(ring)
    for x, y, d, color in ((Inches(10.9), Inches(5.2), Inches(1.3), CYAN), (Inches(12.2), Inches(4.4), Inches(0.6), ICE),
                           (Inches(9.6), Inches(6.3), Inches(0.5), WHITE)):
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, d, d)
        solid(dot, color)
        no_line(dot)
    star = slide.shapes.add_shape(MSO_SHAPE.STAR_4_POINT, Inches(8.9), Inches(5.5), Inches(0.9), Inches(0.9))
    solid(star, WHITE)
    no_line(star)

    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.7), Inches(3.4), Inches(0.55))
    pill.adjustments[0] = 0.5
    solid(pill, WHITE)
    no_line(pill)
    fill_text(pill, t("PROPOSAL · NOV 2026", "企画書 · 2026年11月"), 15, BLUE)
    text(slide, Inches(0.9), Inches(2.55), Inches(8.5), Inches(2.4),
         [t("Kivotos", "キヴォトス"), t("Joint Festival", "合同祭")], 66, WHITE, bold=True)
    text(slide, Inches(0.9), Inches(5.0), Inches(7.5), Inches(1.0),
         t("Ten schools, one weekend, and a lot of fireworks.", "10の学園が集う、花火あふれる週末。"),
         22, RGBColor(0xE6, 0xF4, 0xFF))
    text(slide, Inches(0.9), Inches(6.45), Inches(7.5), Inches(0.5),
         t("Presented by Schale to the General Student Council", "シャーレ → 連邦生徒会"), 15, RGBColor(0xCF, 0xE8, 0xFF))
    notes(slide, t("Open with the halo on screen and wait for the room to settle.",
                   "ヘイローを映したまま、会場が落ち着くのを待ってから始める。"))
    return slide


def pillars_slide():
    slide = prs.slides.add_slide(BLANK)
    background(slide, PAPER)
    heading(slide, t("Why do it together?", "なぜ合同で開催するのか"), t("THE IDEA", "コンセプト"))
    cards = [
        (DEEP, BLUE, "1", t("Every school", "全学園が参加"),
         t("Abydos to Wildhunt, each with a booth of its own.", "アビドスからワイルドハントまで、各校がブースを出展。")),
        (BLUE, CYAN, "2", t("One weekend", "週末の2日間"),
         t("Two days in the main square and the old station.", "中央広場と旧駅舎で2日間開催。")),
        (NAVY, DEEP, "3", t("Shared costs", "費用を分担"),
         t("Split ten ways, the budget goes three times as far.", "10校で分担すれば、予算は3倍に。")),
    ]
    shapes = []
    for i, (color, end, number, title, body) in enumerate(cards):
        x = Inches(0.8 + i * 4.0)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.5), Inches(3.6), Inches(3.7))
        card.adjustments[0] = 0.08
        gradient(card, color, end, angle=45)
        no_line(card)
        shadow(card)
        badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.4), Inches(2.9), Inches(0.9), Inches(0.9))
        solid(badge, WHITE)
        no_line(badge)
        fill_text(badge, number, 26, color)
        text(slide, x + Inches(0.4), Inches(4.1), Inches(2.9), Inches(0.6), title, 24, WHITE, bold=True)
        text(slide, x + Inches(0.4), Inches(4.8), Inches(2.9), Inches(1.6), body, 18, RGBColor(0xEE, 0xF7, 0xFF))
        shapes.append(card)
    notes(slide, t("Yuuka will ask about the budget. Point her to the shared costs.",
                   "ユウカが予算について質問するはず。費用分担の話をする。"))
    return slide, shapes


def chart_slide():
    slide = prs.slides.add_slide(BLANK)
    background(slide, PAPER)
    heading(slide, t("Visitors we expect", "来場者の見込み"), t("ATTENDANCE", "来場者数"), DEEP)
    data = CategoryChartData()
    data.categories = [t(en, ja) for en, ja in (("Millennium", "ミレニアム"), ("Gehenna", "ゲヘナ"), ("Trinity", "トリニティ"),
                                                ("Abydos", "アビドス"), ("Hyakkiyako", "百鬼夜行"), ("Valkyrie", "ヴァルキューレ"))]
    data.add_series(t("Last year", "昨年"), (3200, 4100, 3800, 600, 2100, 1500))
    data.add_series(t("This year", "今年"), (4600, 5200, 4900, 1400, 2900, 2200))
    frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.8), Inches(2.3), Inches(11.7), Inches(4.8), data)
    chart = frame.chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.TOP
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(14)
    chart.legend.font.name = FONT
    for series, color in zip(chart.series, (ICE, BLUE)):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = color
    chart.plots[0].gap_width = 60
    chart.plots[0].overlap = -10
    for axis in (chart.category_axis, chart.value_axis):
        axis.tick_labels.font.size = Pt(13)
        axis.tick_labels.font.name = FONT
        axis.tick_labels.font.color.rgb = MUTED
        axis.format.line.color.rgb = RGBColor(0xD3, 0xE2, 0xF5)
    chart.value_axis.major_gridlines.format.line.color.rgb = RGBColor(0xE3, 0xED, 0xF9)
    chart.value_axis.tick_labels.number_format = "#,##0"
    chart.value_axis.tick_labels.number_format_is_linked = False
    notes(slide, t("Millennium's numbers come from Noa, so they are probably exact.",
                   "ミレニアムの数字はノア調べなので、たぶん正確。"))
    return slide, frame


def timeline_slide():
    slide = prs.slides.add_slide(BLANK)
    background(slide, PAPER)
    heading(slide, t("The road to opening day", "開幕までの道のり"), t("TIMELINE", "スケジュール"))
    stages = [
        (NAVY, t("Plan", "企画"), t("September", "9月"), t("Budget signed off by Yuuka", "ユウカが予算を承認")),
        (DEEP, t("Build", "設営"), t("October", "10月"), t("Stages by the Engineering Club", "エンジニア部がステージを設営")),
        (BLUE, t("Rehearse", "リハーサル"), t("Early November", "11月上旬"), t("A full run with every school", "全学園で通し稽古")),
        (CYAN, t("Open", "開幕"), t("14 November", "11月14日"), t("Fireworks at eight", "20時に花火")),
    ]
    steps = []
    for i, (color, title, when, body) in enumerate(stages):
        x = Inches(0.8 + i * 3.0)
        chevron = slide.shapes.add_shape(MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON,
                                         x, Inches(2.7), Inches(3.1), Inches(1.1))
        solid(chevron, color)
        no_line(chevron)
        fill_text(chevron, title, 22)
        chevron.name = title
        details = text(slide, x, Inches(4.2), Inches(2.9), Inches(1.9),
                       [(when, 18, color, True), (body, 16, MUTED, False)], 16, spacing=Pt(6))
        details.name = when
        # Room on the left for the order badge Dazzle shows while animating.
        details.text_frame.margin_left = Inches(0.55)
        steps.append([(chevron, "fly"), (details, "fade")])
    animate(slide, steps)
    notes(slide, t("Tap through one stage at a time.", "ステージごとにタップして進める。"))
    return slide


def budget_slide():
    slide = prs.slides.add_slide(BLANK)
    background(slide, PAPER)
    heading(slide, t("Where the budget goes", "予算の内訳"), t("BUDGET", "予算"), DEEP)
    data = CategoryChartData()
    items = [(t("Stages", "ステージ"), 38), (t("Food stalls", "屋台"), 22), (t("Fireworks", "花火"), 18),
             (t("Security", "警備"), 12), (t("Ferries", "連絡船"), 10)]
    data.categories = [name for name, _ in items]
    data.add_series(t("Share", "割合"), [share / 100 for _, share in items])
    frame = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(0.6), Inches(2.2), Inches(5.6), Inches(5.0), data)
    chart = frame.chart
    chart.has_legend = False
    plot = chart.plots[0]
    for point, color in zip(plot.series[0].points, (NAVY, DEEP, BLUE, CYAN, ICE)):
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = color
    rows = [(t("Item", "項目"), t("Credits", "クレジット"), t("Paid by", "負担"))] + [
        (t("Stages", "ステージ"), "1,140,000", t("Millennium", "ミレニアム")),
        (t("Food stalls", "屋台"), "660,000", t("Gehenna", "ゲヘナ")),
        (t("Fireworks", "花火"), "540,000", t("Problem Solver 68", "便利屋68")),
        (t("Security", "警備"), "360,000", t("Valkyrie", "ヴァルキューレ")),
        (t("Ferries", "連絡船"), "300,000", t("Odyssey", "オデュッセイア")),
        (t("Total", "合計"), "3,000,000", ""),
    ]
    table = slide.shapes.add_table(len(rows), 3, Inches(6.6), Inches(2.5), Inches(6.0), Inches(4.2)).table
    table.columns[0].width = Inches(2.0)
    table.columns[1].width = Inches(1.8)
    table.columns[2].width = Inches(2.2)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            paragraph = cell.text_frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.RIGHT if c == 1 else PP_ALIGN.LEFT
            font = paragraph.runs[0].font if paragraph.runs else paragraph.font
            font.name = FONT
            font.size = Pt(15)
            font.bold = r in (0, len(rows) - 1)
            font.color.rgb = WHITE if r == 0 else INK
            cell.fill.solid()
            cell.fill.fore_color.rgb = BLUE if r == 0 else (
                RGBColor(0xD6, 0xEB, 0xFF) if r == len(rows) - 1 else (WHITE if r % 2 else RGBColor(0xEE, 0xF6, 0xFF)))
    notes(slide, t("Problem Solver 68 offered to pay for the fireworks. Get it in writing.",
                   "便利屋68が花火代を出すと言っている。書面でもらうこと。"))
    return slide, frame


def closing_slide():
    slide = prs.slides.add_slide(BLANK)
    background(slide, BLUE, NAVY)
    for x, y, d, color in ((Inches(-1.2), Inches(4.4), Inches(4.6), CYAN), (Inches(10.6), Inches(-1.4), Inches(4.2), ICE)):
        ring = slide.shapes.add_shape(MSO_SHAPE.DONUT, x, y, d, d)
        ring.adjustments[0] = 0.1
        solid(ring, color)
        no_line(ring)
    text(slide, Inches(1.5), Inches(2.4), Inches(10.3), Inches(1.4),
         t("See you at the festival!", "合同祭で会いましょう！"), 54, WHITE, bold=True, align=PP_ALIGN.CENTER)
    text(slide, Inches(1.5), Inches(3.9), Inches(10.3), Inches(0.8),
         t("Questions to Schale, any time before 31 October.", "ご質問は10月31日までにシャーレへ。"),
         22, RGBColor(0xDD, 0xEE, 0xFF), align=PP_ALIGN.CENTER)
    return slide


title_slide()
pillars_slide()
chart_slide()
timeline_slide()
budget_slide()
closing_slide()


# MARK: - Saving

def without_docprops(data):
    return re.sub(rb"<(Override|Relationship)\b[^>]*docProps/[^>]*/>", b"", data)


def save(name):
    """Writes the deck without docProps/, which a deck saved by Dazzle has none of."""
    path = os.path.join(OUT, name)
    prs.save(path)
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            if item.filename.startswith("docProps/"):
                continue
            data = src.read(item.filename)
            if item.filename in ("[Content_Types].xml", "_rels/.rels"):
                data = without_docprops(data)
            dst.writestr(item, data)
    os.replace(tmp, path)


save(t("Festival Proposal.pptx", "合同祭のご提案.pptx"))
print(OUT)
